"""ATS lists, the Mode A / Mode B scoring rules, layout facts, the before score and the dashboard."""
from datetime import date

import pytest
from sqlalchemy import select

from app.models.cv import CV
from app.models.unmatched import UnmatchedTerm
from app.services import dashboard_service, parsing_service
from app.services.ats import facts as facts_mod
from app.services.ats import layout, lists, score
from app.services.ats import text as tx
from app.services.pdf_service import PdfError, get_pdf_renderer
from app.services.storage_service import LocalStorage
from tests.conftest import PDF_MIME
from tests.test_parsing import FakeParser, make_text_docx, make_text_pdf
from tests.test_download_export import CONTENT
from tests.test_save_and_events import events

TODAY = date(2026, 10, 29)
GOOD_LAYOUT = {"text": "ok", "headings": 3, "tables": 0, "textboxes": 0, "icons": 0, "pages": 1}


def cv_content(**over):
    content = {
        "personal": {"name": "A B", "email": "a@b.co", "phone": "1", "location": "Pune", "linkedin": "in/ab",
                     "headline": "Engineer", "portfolio": ""},
        "summary": "Engineer.",
        "experience": [{"title": "Senior Engineer", "company": "Z", "start": "2018-01", "end": "", "current": "true",
                        "description": "• Led a team of 5 engineers to cut costs by 20% across three product lines.\n"
                                       "• Built a monitoring platform that reduced incident response time by 35%."}],
        "education": [{"start": "2010-06", "end": "2014-05"}],
        "skills": ["Python", "K8s", "SQL", "Go", "Docker"],
    }
    content.update(over)
    return content


def check(result, cid):
    return next(c for c in result.checks if c["id"] == cid)


# --- lists -------------------------------------------------------------------------------------

def test_no_word_is_in_two_lists():
    assert not set(lists.ACTION_VERBS) & set(lists.WEAK_OPENERS)
    forms = [f for c, others in lists.ALIASES.items() for f in (c, *others)]
    assert len(forms) == len(set(forms)), "an alias form belongs to two skills"
    assert len(lists.ACTION_VERBS) == len(set(lists.ACTION_VERBS))
    for family, rungs in lists.LADDERS.items():
        assert len(rungs) == len(set(rungs)), family


def test_mode_points_add_up():
    assert sum(score.MODE_A.values()) == 100
    assert sum(score.MODE_B_CV.values()) == 60 and sum(score.MODE_B_JOB.values()) == 40


@pytest.mark.parametrize("word", ["managed", "manages", "managing", "led", "lead", "built", "build", "planned",
                                  "plan", "optimised", "created", "cut", "drove", "grew"])
def test_strong_verbs_match_in_any_form(word):
    assert tx.is_strong_verb(word)


@pytest.mark.parametrize("word", ["responsible", "assisted", "helped", "the", "team"])
def test_weak_words_are_not_strong(word):
    assert not tx.is_strong_verb(word)


def test_metrics_ignore_years_used_as_dates():
    assert tx.has_metric("Cut costs by 20%") and tx.has_metric("Managed a team of 5") and tx.has_metric("Saved ₹ lakh")
    assert not tx.has_metric("Joined in March 2021 and led the team")


def test_skill_matching_uses_aliases_and_whole_words():
    assert tx.skill_in("deployed on k8s clusters", "Kubernetes")
    assert tx.skill_in("worked with node.js", "NodeJS")
    assert not tx.skill_in("reactive programming", "React")


# --- Mode A ------------------------------------------------------------------------------------

def test_a_complete_cv_scores_high_and_has_no_fix_for_passed_checks():
    result = score.score_mode_a(facts_mod.from_content(cv_content()), GOOD_LAYOUT, TODAY)
    assert result.score >= 85
    assert {c["id"] for c in result.checks} == set(score.MODE_A)
    assert all(c["fix"] is None for c in result.checks if c["passed"])
    assert result.public()["framework"] == "ats-framework-v2"


def test_checks_use_the_framework_points():
    f = facts_mod.from_content(cv_content())
    result = score.score_mode_a(f, GOOD_LAYOUT, TODAY)
    assert check(result, "A3")["points"] == 11 and check(result, "A1")["max"] == 12
    # mixed date formats (month+year jobs, month+year education here) are consistent
    assert check(result, "A5")["points"] == 9
    mixed = cv_content(education=[{"start": "2010", "end": "2014"}])
    assert check(score.score_mode_a(facts_mod.from_content(mixed), GOOD_LAYOUT, TODAY), "A5")["points"] == 4


def test_contact_points_follow_the_split():
    c = cv_content()
    c["personal"].update(phone="", linkedin="")
    result = score.score_mode_a(facts_mod.from_content(c), GOOD_LAYOUT, TODAY)
    assert check(result, "A3")["points"] == 7  # name 2 + email 3 + location 2
    assert "phone" in check(result, "A3")["fix"] and "LinkedIn" in check(result, "A3")["fix"]


def test_verb_metric_filler_and_length_checks():
    desc = ("• Responsible for the weekly reports\n• Worked on dashboards\n• Helped the team with forecasting\n"
            "• Assisted in audits\n• Dealt with vendors")
    job = {"title": "T", "company": "C", "start": "2020-01", "end": "2021-01", "description": desc}
    result = score.score_mode_a(facts_mod.from_content(cv_content(experience=[job])), GOOD_LAYOUT, TODAY)
    assert check(result, "A6")["points"] == 0       # no strong verbs
    assert check(result, "A7")["points"] == 0       # no numbers
    assert check(result, "A8")["points"] == 0       # more than 3 filler phrases
    assert check(result, "A9")["points"] == 0       # all short
    assert result.first_fix["check_id"] == "A6"     # biggest loss, ties go to the earlier check


def test_a_verb_that_opens_more_than_four_bullets_fails_a10():
    desc = "\n".join(f"• Managed project number {i} with a group of ten people over a long period of time"
                     for i in range(5))
    job = {"title": "T", "company": "C", "start": "2020-01", "end": "2021-01", "description": desc}
    f = facts_mod.from_content(cv_content(experience=[job]))
    assert check(score.score_mode_a(f, GOOD_LAYOUT, TODAY), "A10")["points"] == 0


def test_layout_checks_are_left_out_when_there_is_no_layout():
    result = score.score_mode_a(facts_mod.from_content(cv_content()), None, TODAY)
    assert {c["id"] for c in result.checks}.isdisjoint({"A1", "A2", "A4", "A11"})
    assert 0 < result.score <= 100  # earned / applicable, not zero for the missing checks


def test_layout_traps_and_length():
    f = facts_mod.from_content(cv_content())
    trap = {**GOOD_LAYOUT, "tables": 2, "textboxes": 1, "icons": 3}
    assert check(score.score_mode_a(f, trap, TODAY), "A4")["points"] == 0      # 8 - 4 - 2 - 3, floor 0
    assert check(score.score_mode_a(f, {**GOOD_LAYOUT, "pages": 3}, TODAY), "A11")["points"] == 3  # one band off
    assert check(score.score_mode_a(f, {**GOOD_LAYOUT, "pages": 4}, TODAY), "A11")["points"] == 0  # 8 years: up to 2
    assert check(score.score_mode_a(f, {**GOOD_LAYOUT, "pages": 2}, TODAY), "A11")["points"] == 6
    assert check(score.score_mode_a(f, {**GOOD_LAYOUT, "text": "unreadable"}, TODAY), "A1")["points"] == 0


def test_unlisted_past_tense_verbs_are_returned_for_review():
    job = {"title": "T", "company": "C", "start": "2020-01", "end": "", "current": "true",
           "description": "• Wrangled messy data sets every week"}
    f = facts_mod.from_content(cv_content(experience=[job]))
    assert ("verb", "wrangled") in score.score_mode_a(f, GOOD_LAYOUT, TODAY).unmatched


# --- Mode B (the scoring rules; the Align to a job screen is not built yet) -------------------

def test_mode_b_scales_not_applicable_checks_out():
    f = facts_mod.from_content(cv_content())
    req = score.Requirements(job_title="Lead Engineer", core_skills=["Kubernetes", "Go"],
                             secondary_skills=["Python"], years_required=5, needs_management=True,
                             sector="Software")
    result = score.score_mode_b(f, GOOD_LAYOUT, req, TODAY)
    assert 0 <= result.score <= 100
    assert check(result, "B1")["points"] == 14.0   # both found (K8s alias, Go)
    assert check(result, "B5")["points"] == 6      # Senior Engineer is the direct precursor of Lead Engineer
    none = score.score_mode_b(f, GOOD_LAYOUT, score.Requirements(job_title="Zookeeper"), TODAY)
    assert {c["id"] for c in none.checks}.isdisjoint({"B1", "B2", "B3", "B4", "B5", "B6", "B7"})
    assert ("title", "Zookeeper") in none.unmatched


# --- layout facts ------------------------------------------------------------------------------

def test_text_integrity_levels():
    assert layout.text_integrity("short") == "unreadable"
    assert layout.text_integrity("Experience at Acme. " * 20) == "ok"
    assert layout.text_integrity("�" * 30 + "readable words here " * 20) == "scrambled"


def test_pdf_layout_counts_pages_and_headings():
    pdf = make_text_pdf("Jane Doe\nProfessional experience\nAcme Ltd, Data Analyst, built dashboards for sales. " * 3
                        + "\nEducation\nUniversity of Leeds, BSc Mathematics\nSkills\nPython, SQL, Tableau" * 2)
    facts = layout.from_pdf(pdf)
    assert facts["pages"] == 1 and facts["headings"] == 3 and facts["text"] == "ok" and facts["tables"] == 0


def test_docx_layout_counts_tables():
    data = make_text_docx(["Experience " * 30], table_rows=[["a", "b"], ["c", "d"]])
    facts = layout.from_docx(data)
    assert facts["tables"] == 1 and facts["textboxes"] == 0


def test_unreadable_file_has_no_layout():
    assert layout.from_pdf(b"%PDF-1.4 garbage") is None


# --- CV strength -------------------------------------------------------------------------------

def test_cv_strength_points_and_next_step():
    full = dashboard_service.cv_strength(cv_content(projects=[{"id": "p"}], certifications=[{"name": "x"}]))
    assert full["percent"] == 100 and full["next_step"] is None
    part = dashboard_service.cv_strength(cv_content())
    assert part["percent"] == 80 and part["next_step"]["section"] in ("projects", "others")
    assert part["next_step"]["points"] == 10 and part["next_step"]["text"].endswith("to reach 90%.")
    assert dashboard_service.cv_strength(None)["percent"] == 0


def test_a_role_without_dates_is_not_complete():
    c = cv_content(experience=[{"title": "T", "company": "C", "start": "", "end": ""}])
    experience = next(s for s in dashboard_service.cv_strength(c)["sections"] if s["key"] == "experience")
    assert experience["points"] == 10 and not experience["complete"]


# --- API ---------------------------------------------------------------------------------------

class Renderer:
    def __init__(self, pdf=None, fail=False):
        self.pdf, self.fail, self.calls = pdf, fail, 0

    def render(self, draft):
        self.calls += 1
        if self.fail:
            raise PdfError
        return self.pdf


GOOD_PDF = make_text_pdf("Jane Doe\nProfessional experience\nAcme Ltd, Data Analyst, built dashboards for sales. "
                         "\nEducation\nUniversity of Leeds, BSc Mathematics, 2014 to 2017, first class honours degree\n"
                         "Skills\nPython, SQL, Tableau, statistics, communication, stakeholder management")


@pytest.fixture()
def renderer(client):
    r = Renderer(GOOD_PDF)
    client.app.dependency_overrides[get_pdf_renderer] = lambda: r
    return r


def test_dashboard_without_a_cv_shows_the_upload_state(client, auth):
    body = client.get("/api/dashboard", headers=auth[0]).json()
    assert body == {"state": "no_cv", "cv_strength": None, "ats": None, "plan": "free"}


def test_dashboard_needs_login(client):
    assert client.get("/api/dashboard").status_code == 401
    assert client.get("/api/ats").status_code == 401


def test_dashboard_while_parsing_and_when_parsing_failed(client, auth):
    client.post("/api/cv/upload", headers=auth[0], files={"file": ("cv.pdf", make_text_pdf("x"), PDF_MIME)})
    assert client.get("/api/dashboard", headers=auth[0]).json()["state"] == "parsing"
    with client.session_factory() as db:
        cv = db.scalar(select(CV))
        cv.parsing_status = "failed"
        db.commit()
    body = client.get("/api/dashboard", headers=auth[0]).json()
    assert body["state"] == "needs_details" and body["cv_strength"] is None and body["ats"] is None


def test_dashboard_scores_the_saved_cv(client, auth, renderer):
    assert client.patch("/api/cv", headers=auth[0], json={"content": CONTENT}).status_code == 200
    body = client.get("/api/dashboard", headers=auth[0]).json()
    assert body["state"] == "ready" and body["plan"] == "free"
    assert 0 < body["cv_strength"]["percent"] <= 100
    assert body["ats"]["framework"] == "ats-framework-v2" and 0 < body["ats"]["score"] <= 100
    ids = {c["id"] for c in body["ats"]["checks"]}
    assert {"A1", "A2", "A4", "A11"} <= ids                 # layout read from the generated PDF
    assert body["ats"]["first_fix"]["text"]


def test_the_generated_pdf_is_only_rendered_when_the_content_changes(client, auth, renderer):
    client.patch("/api/cv", headers=auth[0], json={"content": CONTENT})
    client.get("/api/dashboard", headers=auth[0])
    client.get("/api/dashboard", headers=auth[0])
    client.get("/api/ats", headers=auth[0])
    assert renderer.calls == 1
    client.patch("/api/cv", headers=auth[0], json={"content": {**CONTENT, "summary": "Changed."}})
    client.get("/api/ats", headers=auth[0])
    assert renderer.calls == 2


def test_a_failed_render_leaves_the_layout_checks_out(client, auth):
    r = Renderer(fail=True)
    client.app.dependency_overrides[get_pdf_renderer] = lambda: r
    client.patch("/api/cv", headers=auth[0], json={"content": CONTENT})
    body = client.get("/api/dashboard", headers=auth[0]).json()
    assert body["ats"]["score"] > 0
    assert {c["id"] for c in body["ats"]["checks"]}.isdisjoint({"A1", "A2", "A4", "A11"})


def test_ats_endpoint_has_no_before_score_without_an_upload(client, auth, renderer):
    client.patch("/api/cv", headers=auth[0], json={"content": CONTENT})
    body = client.get("/api/ats", headers=auth[0]).json()
    assert body["before"] is None and body["after"]["score"] > 0


def test_ats_endpoint_without_a_cv(client, auth):
    assert client.get("/api/ats", headers=auth[0]).json() == {"before": None, "after": None}


def test_parsing_keeps_the_before_score_and_records_unmatched_terms(client, auth, renderer):
    storage = LocalStorage(str(client.upload_dir))
    client.post("/api/cv/upload", headers=auth[0], files={"file": ("cv.pdf", GOOD_PDF, PDF_MIME)})
    cv_id, key = client.queued_parse_jobs[-1]
    with client.session_factory() as db:
        parsing_service.run_parsing(db, storage, FakeParser(), cv_id, key)
        cv = db.get(CV, cv_id)
        before = cv.ats_before
        assert before["framework"] == "ats-framework-v2" and before["layout"]["headings"] == 3
        assert "layout" in before and before["checks"]
        assert ({t.kind for t in db.scalars(select(UnmatchedTerm))} - {"verb", "skill", "title"}) == set()
    assert not list(client.upload_dir.rglob("*.pdf"))          # the file itself is still deleted
    body = client.get("/api/ats", headers=auth[0]).json()
    assert body["before"]["score"] == before["score"] and "layout" not in body["before"]
    # no saved canvas yet: the current score is the parsed upload's
    assert body["after"]["score"] > 0


def test_a_new_upload_clears_the_old_scores(client, auth, renderer):
    storage = LocalStorage(str(client.upload_dir))
    client.post("/api/cv/upload", headers=auth[0], files={"file": ("cv.pdf", GOOD_PDF, PDF_MIME)})
    cv_id, key = client.queued_parse_jobs[-1]
    with client.session_factory() as db:
        parsing_service.run_parsing(db, storage, FakeParser(), cv_id, key)
    client.post("/api/cv/upload", headers=auth[0], files={"file": ("cv2.pdf", GOOD_PDF, PDF_MIME)})
    with client.session_factory() as db:
        assert db.get(CV, cv_id).ats_before is None


def test_next_viewed_is_recorded_once_per_visit_without_content(client, auth):
    assert client.post("/api/next/viewed", headers=auth[0]).status_code == 204
    assert events(client, "next_viewed") == [("next_viewed", {})]


def test_cli_sets_a_users_plan_for_testing(client, auth, monkeypatch):
    from app import cli
    monkeypatch.setattr(cli, "SessionLocal", client.session_factory)
    assert cli.main(["users", "set-plan", "USER@example.com", "pro"]) == 0
    assert client.get("/api/dashboard", headers=auth[0]).json()["plan"] == "pro"
    assert cli.main(["users", "set-plan", "nobody@example.com", "pro"]) == 1


def test_a_blank_job_entry_is_ignored():
    content = cv_content()
    content["experience"].append({"id": "e2", "title": "", "company": "", "start": "", "end": "", "description": ""})
    result = score.score_mode_a(facts_mod.from_content(content), GOOD_LAYOUT, TODAY)
    assert check(result, "A5")["points"] == 9
    assert len(facts_mod.from_content(content).roles) == 1


def test_the_house_style_names_exactly_the_ats_action_verbs():
    """Rule 15 of the house style lists the verbs the A6 check accepts: the two must not drift apart."""
    from pathlib import Path
    text = (Path(__file__).resolve().parents[1] / "app" / "prompts" / "cv_house_style.md").read_text()
    line = text.split("one of these verbs when it\n    says the same thing: ", 1)[1].split(". Keep the meaning", 1)[0]
    named = {v.strip().lower() for v in line.replace("\n", " ").split(",")}
    assert named == set(lists.ACTION_VERBS)


def test_the_header_experience_must_be_work_experience_professional_experience_or_work_history():
    cases = (("Professional experience", 3), ("Work experience", 3), ("Work history", 3), ("Experience", 2))
    for heading, expected in cases:
        text = f"{heading}\nAcme\nEducation\nLeeds\nSkills\nPython\n"
        assert layout.headings_found(text) == expected, heading


def test_a_role_needs_a_job_title_and_an_employer():
    def role(title, company):
        return {"title": title, "company": company, "start": "2020-01", "end": "2021-01", "description": ""}

    def a12_check(*roles):
        f = facts_mod.from_content(cv_content(experience=list(roles)))
        return check(score.score_mode_a(f, GOOD_LAYOUT, TODAY), "A12")

    def a12(*roles):
        return a12_check(*roles)["points"]
    assert a12(role("Analyst", "Acme"), role("Intern", "Beta")) == 6
    assert a12(role("Analyst", "Acme"), role("", "Beta")) == 3            # half of the roles are complete
    assert a12(role("Analyst", ""), role("", "Beta"), role("Intern", "Gamma")) == 0
    assert a12() == 0
    fix = a12_check(role("", "Beta"))["fix"]
    assert "job title" in fix and "employer" in fix


def test_contact_details_only_in_a_word_header_cost_points_in_the_layout_check():
    import io
    import docx

    def make(header_text):
        document = docx.Document()
        document.add_paragraph("Professional experience " * 20)
        document.sections[0].header.paragraphs[0].text = header_text
        buf = io.BytesIO()
        document.save(buf)
        return layout.from_docx(buf.getvalue())

    only_header = make("Jane Doe | jane@example.com | +44 7700 900123")
    assert only_header["header_contact"] is True
    f = facts_mod.from_content(cv_content())
    result = score.score_mode_a(f, {**GOOD_LAYOUT, "header_contact": True}, TODAY)
    assert check(result, "A4")["points"] == 6 and "header" in check(result, "A4")["fix"]
    assert make("Jane Doe")["header_contact"] is False
    # an email that is also in the page body is read by every parser
    document = docx.Document()
    document.add_paragraph("jane@example.com " + "Professional experience " * 20)
    document.sections[0].header.paragraphs[0].text = "jane@example.com"
    buf = io.BytesIO()
    document.save(buf)
    assert layout.from_docx(buf.getvalue())["header_contact"] is False


def test_checks_belong_to_groups_and_each_group_adds_up():
    result = score.score_mode_a(facts_mod.from_content(cv_content()), GOOD_LAYOUT, TODAY)
    groups = {g["key"]: g for g in result.groups}
    assert list(groups) == ["searchability", "writing"]
    assert groups["searchability"]["max"] == 12 + 12 + 11 + 8 + 9 + 6 + 6 and groups["writing"]["max"] == 36
    for key, g in groups.items():
        members = [c for c in result.checks if c["group"] == key]
        assert g["points"] == round(sum(c["points"] for c in members), 1)
        assert g["issues"] == sum(1 for c in members if c["fix"])
        assert g["percent"] == round(g["points"] / g["max"] * 100)
    assert sum(g["max"] for g in groups.values()) == 100


def test_job_match_has_four_groups_and_a_group_with_nothing_applicable_is_left_out():
    f = facts_mod.from_content(cv_content())
    req = score.Requirements(job_title="Lead Engineer", core_skills=["Kubernetes"], years_required=5)
    result = score.score_mode_b(f, GOOD_LAYOUT, req, TODAY)
    assert [g["key"] for g in result.groups] == ["searchability", "writing", "hard_skills", "job_fit"]
    bare = score.score_mode_b(f, GOOD_LAYOUT, score.Requirements(job_title="Zookeeper"), TODAY)
    assert [g["key"] for g in bare.groups] == ["searchability", "writing"]


def test_failing_checks_say_what_the_canvas_can_do_about_them():
    mixed = cv_content(education=[{"start": "2010", "end": "2014"}])
    result = score.score_mode_a(facts_mod.from_content(mixed), {**GOOD_LAYOUT, "pages": 3}, TODAY)
    kinds = {c["id"]: c["fix_kind"] for c in result.checks}
    assert kinds["A5"] == "dates_format"                       # mixed formats: fixed automatically
    assert kinds["A8"] is None and kinds["A3"] is None         # passing checks offer nothing
    missing_dates = cv_content(experience=[{"title": "T", "company": "C", "start": "", "end": "", "description": ""}])
    assert {c["id"]: c["fix_kind"] for c in score.score_mode_a(
        facts_mod.from_content(missing_dates), GOOD_LAYOUT, TODAY).checks}["A5"] == "go_experience"
    c = cv_content()
    c["personal"]["phone"] = ""
    assert {x["id"]: x["fix_kind"] for x in score.score_mode_a(
        facts_mod.from_content(c), GOOD_LAYOUT, TODAY).checks}["A3"] == "go_contact"
    # layout problems cannot be fixed on the Canvas
    assert {x["id"]: x["fix_kind"] for x in result.checks}["A11"] is None


def test_the_api_returns_groups_and_fix_kinds(client, auth, renderer):
    client.patch("/api/cv", headers=auth[0], json={"content": CONTENT})
    body = client.get("/api/dashboard", headers=auth[0]).json()["ats"]
    assert [g["key"] for g in body["groups"]] == ["searchability", "writing"]
    assert all("group" in c and "fix_kind" in c for c in body["checks"])


def _bullets_content(bullets):
    job = {"title": "T", "company": "C", "start": "2020-01", "end": "2021-01",
           "description": "\n".join(f"• {b}" for b in bullets)}
    return facts_mod.from_content(cv_content(experience=[job]))


def test_a6_names_the_words_that_count_against_the_user():
    f = _bullets_content(["Liaised with vendors about onboarding for new joiners in four regional offices each year",
                          "Compiled monthly headcount reports for the finance director and the board",
                          "Led a team of five engineers to cut costs by 20% across three product lines"])
    a6 = check(score.score_mode_a(f, GOOD_LAYOUT, TODAY), "A6")
    assert "(Liaised, Compiled)" in a6["fix"]          # only the words that count against them
    assert a6["fix_kind"] == "improve_wording"          # the suggestions can swap in a listed verb


def test_a6_when_only_weak_openers_remain_it_says_the_user_has_to_reword_them():
    f = _bullets_content(["Supported payroll audits across four regional offices every quarter",
                          "Helped with recruitment drives for store roles in three cities",
                          "Assisted the HR manager with employee relations cases"])
    a6 = check(score.score_mode_a(f, GOOD_LAYOUT, TODAY), "A6")
    assert "Supported" in a6["fix"] and "left as written" in a6["fix"]
    assert a6["fix_kind"] == "go_experience"             # Improve all would not change these


def test_common_strong_verbs_added_on_2026_10_31_count_as_strong():
    for word in ("coordinated", "conducted", "prepared", "oversaw", "wrote", "spearheaded", "organised", "onboarded"):
        assert tx.is_strong_verb(word), word
    for form in ("oversee", "overseeing", "writing", "coordinates", "preparing"):
        assert tx.is_strong_verb(form), form
    for weak in ("supported", "helped", "assisted", "handled", "responsible"):
        assert not tx.is_strong_verb(weak), weak
