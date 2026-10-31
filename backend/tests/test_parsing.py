import io
import json
from types import SimpleNamespace

import docx
import pytest

from app.models.cv import CV
from app.services import parsing_service
from app.services.cv_parser import (
    PARTS, BasicsPart, ClaudeCvParser, ExperiencePart, ExtraSections, ParsedResume, ParseError, merge_fragment,
)
from app.services.storage_service import LocalStorage
from app.services.text_extraction import ExtractionError, extract_text
from tests.conftest import DOCX_MIME, PDF_MIME, make_pdf

CV_TEXT = (
    "Jane Doe - Senior Data Analyst - jane@example.com\n"
    "Experience: Acme Ltd, Data Analyst, 2019-03 to Present. Built weekly sales dashboards.\n"
    "Education: University of Leeds, BSc Mathematics, 2015 - 2018.\n"
    "Skills: Python, SQL, Tableau, statistics, stakeholder communication."
)


def parts(contact="complete", basics="complete", experience="complete", extras="complete"):
    return {"contact": contact, "basics": basics, "experience": experience, "extras": extras}


PARSED = {"basics": {"name": "Jane Doe"}, "work": [{"name": "Acme Ltd", "position": "Data Analyst"}]}


def make_text_pdf(text: str) -> bytes:
    """A minimal, valid one-page PDF whose text layer contains `text`."""
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    lines = escaped.split("\n")
    ops = "BT /F1 10 Tf 40 800 Td 12 TL " + " ".join(f"({line}) Tj T*" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(ops)} >>\nstream\n{ops}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def make_text_docx(paragraphs, table_rows=()) -> bytes:
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def split_parts(resume):
    """Split a full JSON Resume into the fragments the three parsing parts produce."""
    basics = {k: v for k, v in resume.get("basics", {}).items() if k != "summary"}
    experience = {k: resume[k] for k in ("work", "education") if k in resume}
    if "summary" in resume.get("basics", {}):
        experience["basics"] = {"summary": resume["basics"]["summary"]}
    extras = {k: v for k, v in resume.items() if k not in ("basics", "work", "education")}
    return {"basics": {"basics": basics} if basics else {}, "experience": experience, "extras": extras}


def collect(results):
    """Merge a parser's part results the way the pipeline does; returns (resume, failed parts)."""
    resume, errors = {}, {}
    for part, result in results:
        if isinstance(result, ParseError):
            errors[part] = result
        else:
            resume = merge_fragment(resume, result)
    return resume, errors


class FakeParser:
    def __init__(self, result=None, error=None, fail_parts=()):
        self.result, self.error, self.fail_parts = result or PARSED, error, set(fail_parts)
        self.calls = []

    def _parts(self):
        if self.error:
            raise self.error
        for part, fragment in split_parts(self.result).items():
            yield part, (ParseError(f"{part} failed.") if part in self.fail_parts else fragment)

    def parse_text(self, text):
        self.calls.append(("text", text))
        return self._parts()

    def parse_pdf(self, data):
        self.calls.append(("pdf", data))
        return self._parts()


# --- text extraction -------------------------------------------------------------

def test_extract_text_from_pdf_and_docx():
    assert "Senior Data Analyst" in extract_text(make_text_pdf(CV_TEXT), ".pdf")
    text = extract_text(make_text_docx(["Jane Doe", "Data Analyst"], [("Python", "SQL")]), ".docx")
    assert "Jane Doe" in text and "Python | SQL" in text


def test_extract_text_rejects_broken_files():
    with pytest.raises(ExtractionError):
        extract_text(b"%PDF-1.4 garbage", ".pdf")
    with pytest.raises(ExtractionError):
        extract_text(b"PK\x03\x04 not really a zip", ".docx")


# --- pipeline ----------------------------------------------------------------------

@pytest.fixture()
def stored_cv(client, auth):
    """Upload a file through the API and return a helper to run the queued parse job."""
    storage = LocalStorage(str(client.upload_dir))

    def upload_and_run(data, name="cv.pdf", mime=PDF_MIME, parser=None):
        r = client.post("/api/cv/upload", headers=auth[0], files={"file": (name, data, mime)})
        assert r.status_code == 201
        cv_id, key = client.queued_parse_jobs[-1]
        parser = parser or FakeParser()
        with client.session_factory() as db:
            parsing_service.run_parsing(db, storage, parser, cv_id, key)
        return client.get("/api/cv", headers=auth[0]).json(), parser

    return upload_and_run


def test_upload_queues_a_parse_job_without_waiting(client, auth):
    r = client.post("/api/cv/upload", headers=auth[0], files={"file": ("cv.pdf", make_pdf(), PDF_MIME)})
    assert r.status_code == 201 and r.json()["parsing_status"] == "pending"
    assert len(client.queued_parse_jobs) == 1
    cv_id, key = client.queued_parse_jobs[0]
    assert str(cv_id) == r.json()["id"] and key.endswith(".pdf")


def test_text_pdf_is_parsed_and_stored(stored_cv):
    body, parser = stored_cv(make_text_pdf(CV_TEXT))
    assert body["parsing_status"] == "complete" and body["parsing_error"] is None
    # the email comes from the instant pattern step; the (fake) AI parts add the rest
    assert body["json_resume"] == {**PARSED, "basics": {"name": "Jane Doe", "email": "jane@example.com"}}
    assert body["parsing_parts"] == parts()
    assert parser.calls[0][0] == "text" and "Acme Ltd" in parser.calls[0][1]


def test_docx_is_parsed(stored_cv):
    body, parser = stored_cv(make_text_docx(CV_TEXT.split("\n")), "cv.docx", DOCX_MIME)
    assert body["parsing_status"] == "complete" and parser.calls[0][0] == "text"


def test_scanned_pdf_is_sent_to_claude_as_a_file(stored_cv):
    data = make_text_pdf("")  # valid PDF with no text layer, like a scan
    body, parser = stored_cv(data)
    assert body["parsing_status"] == "complete"
    assert parser.calls == [("pdf", data)]


def test_unreadable_pdf_fails_with_message(stored_cv):
    body, parser = stored_cv(make_pdf())  # passes the upload check but is not a real PDF
    assert body["parsing_status"] == "failed" and "could not be read" in body["parsing_error"]
    assert parser.calls == []


def test_docx_without_text_fails_with_message(stored_cv):
    body, parser = stored_cv(make_text_docx([]), "cv.docx", DOCX_MIME)
    assert body["parsing_status"] == "failed" and "No text" in body["parsing_error"]
    assert parser.calls == []


def test_parser_error_message_is_shown(stored_cv):
    body, _ = stored_cv(make_text_pdf(CV_TEXT), parser=FakeParser(error=ParseError("Service down.")))
    assert body["parsing_status"] == "failed" and body["parsing_error"] == "Service down."
    # the instant contact step still saved what it found, but without an AI part the parse failed
    assert body["json_resume"] == {"basics": {"email": "jane@example.com"}}
    assert body["parsing_parts"] == parts(basics="failed", experience="failed", extras="failed")


def test_a_failed_part_keeps_the_others(stored_cv):
    body, _ = stored_cv(make_text_pdf(CV_TEXT), parser=FakeParser(fail_parts={"experience"}))
    assert body["parsing_status"] == "complete" and body["parsing_error"] is None
    assert body["parsing_parts"] == parts(experience="failed")
    assert body["json_resume"] == {"basics": {"email": "jane@example.com", "name": "Jane Doe"}}


def test_all_parts_failed_is_a_failed_parse(stored_cv):
    body, _ = stored_cv(make_text_pdf(CV_TEXT), parser=FakeParser(fail_parts=set(PARTS)))
    assert body["parsing_status"] == "failed" and body["parsing_error"] == "basics failed."


def test_each_part_is_visible_as_soon_as_it_is_saved(client, auth):
    storage = LocalStorage(str(client.upload_dir))
    client.post("/api/cv/upload", headers=auth[0], files={"file": ("a.pdf", make_text_pdf(CV_TEXT), PDF_MIME)})
    cv_id, key = client.queued_parse_jobs[0]
    seen = []

    class Watching(FakeParser):
        def _parts(self):
            # the AI parts start only after the contact step has been saved
            seen.append(client.get("/api/cv", headers=auth[0]).json())
            for part, fragment in super()._parts():
                yield part, fragment
                seen.append(client.get("/api/cv", headers=auth[0]).json())

    with client.session_factory() as db:
        parsing_service.run_parsing(db, storage, Watching(), cv_id, key)
    # the pattern step comes first: email on screen before any AI part has finished
    first, second = seen[0], seen[1]
    assert first["parsing_parts"] == parts(basics="pending", experience="pending", extras="pending")
    assert first["json_resume"] == {"basics": {"email": "jane@example.com"}}
    assert second["parsing_status"] == "in_progress"
    assert second["parsing_parts"] == parts(experience="pending", extras="pending")
    assert second["json_resume"] == {"basics": {"email": "jane@example.com", "name": "Jane Doe"}}


def test_unexpected_error_shows_generic_message(stored_cv):
    body, _ = stored_cv(make_text_pdf(CV_TEXT), parser=FakeParser(error=RuntimeError("secret detail")))
    assert body["parsing_status"] == "failed"
    assert body["parsing_error"] == parsing_service.GENERIC_FAILURE


def test_job_for_replaced_upload_is_discarded(client, auth):
    storage = LocalStorage(str(client.upload_dir))
    files = {"file": ("a.pdf", make_text_pdf(CV_TEXT), PDF_MIME)}
    client.post("/api/cv/upload", headers=auth[0], files=files)
    old_id, old_key = client.queued_parse_jobs[0]
    client.post("/api/cv/upload", headers=auth[0], files={"file": ("b.pdf", make_text_pdf(CV_TEXT), PDF_MIME)})
    parser = FakeParser()
    with client.session_factory() as db:
        parsing_service.run_parsing(db, storage, parser, old_id, old_key)
    assert parser.calls == []
    assert client.get("/api/cv", headers=auth[0]).json()["parsing_status"] == "pending"


def test_result_discarded_if_cv_replaced_while_parsing(client, auth):
    storage = LocalStorage(str(client.upload_dir))
    client.post("/api/cv/upload", headers=auth[0], files={"file": ("a.pdf", make_text_pdf(CV_TEXT), PDF_MIME)})
    cv_id, key = client.queued_parse_jobs[0]

    class ReplacingParser(FakeParser):
        def parse_text(self, text):
            # Simulate the user uploading a new CV while Claude is working.
            new_file = {"file": ("b.docx", make_text_docx(["x"]), DOCX_MIME)}
            r = client.post("/api/cv/upload", headers=auth[0], files=new_file)
            assert r.status_code == 201, r.text
            return super().parse_text(text)

    with client.session_factory() as db:
        parsing_service.run_parsing(db, storage, ReplacingParser(), cv_id, key)
    body = client.get("/api/cv", headers=auth[0]).json()
    assert body["parsing_status"] == "pending" and body["json_resume"] == {}


def test_queue_failure_marks_cv_failed_but_upload_succeeds(client, auth):
    from app.main import app
    from app.services.parse_queue import get_parse_queue

    def broken(cv_id, key):
        raise ConnectionError("redis down")

    app.dependency_overrides[get_parse_queue] = lambda: broken
    r = client.post("/api/cv/upload", headers=auth[0], files={"file": ("cv.pdf", make_pdf(), PDF_MIME)})
    assert r.status_code == 201
    assert r.json()["parsing_status"] == "failed" and "try uploading it again" in r.json()["parsing_error"]
    assert set(r.json()["parsing_parts"].values()) == {"failed"}


def test_new_upload_resets_previous_result(stored_cv, client, auth):
    stored_cv(make_text_pdf(CV_TEXT))
    r = client.post("/api/cv/upload", headers=auth[0], files={"file": ("cv.pdf", make_pdf(), PDF_MIME)})
    assert r.json()["parsing_status"] == "pending" and r.json()["json_resume"] == {}
    assert set(r.json()["parsing_parts"].values()) == {"pending"}


# --- Claude parser (API client faked; no network) ------------------------------------

class FakeMessages:
    """Answers each request with the part of a full ParsedResume that its schema asks for."""

    def __init__(self, response):
        self.response, self.calls = response, []

    def parse(self, **kwargs):
        self.calls.append(kwargs)
        full, schema = self.response.parsed_output, kwargs["output_format"]
        part = None
        if full is not None:
            data = full.model_dump()
            data["summary"] = data["basics"].get("summary", "")  # used by ExperiencePart only
            part = schema.model_validate(data)
        return SimpleNamespace(**{**vars(self.response), "parsed_output": part})

    @property
    def kwargs(self):
        return self.calls[0]


def fake_client(response):
    messages = FakeMessages(response)
    return SimpleNamespace(beta=SimpleNamespace(messages=messages)), messages


def test_claude_parser_request_and_result():
    output = ParsedResume.model_validate({"basics": {"name": "Jane Doe", "email": ""}, "skills": [{"name": "SQL"}]})
    client, messages = fake_client(SimpleNamespace(stop_reason="end_turn", parsed_output=output, _request_id="r1"))
    result, errors = collect(ClaudeCvParser(client=client, model="claude-sonnet-5-5").parse_text("Jane Doe CV"))
    assert errors == {}
    assert result["basics"] == {"name": "Jane Doe", "profiles": []}
    assert result["skills"] == [{"name": "SQL", "keywords": []}]
    assert {kw["output_format"] for kw in messages.calls} == {BasicsPart, ExperiencePart, ExtraSections}
    for kw in messages.calls:
        assert kw["model"] == "claude-sonnet-5-5"
        assert kw["fallbacks"] == "default" and kw["betas"] == ["server-side-fallback-2026-07-01"]
        assert "Jane Doe CV" in kw["messages"][0]["content"][0]["text"]
        assert "Never invent" in kw["system"][0]["text"]
        assert kw["system"][0]["cache_control"] == {"type": "ephemeral"}
        assert "cache_control" not in kw["system"][1]  # the part's own task is not cached


def test_claude_parser_sends_scanned_pdf_as_document():
    output = ParsedResume.model_validate({"basics": {}})
    client, messages = fake_client(SimpleNamespace(stop_reason="end_turn", parsed_output=output, _request_id="r"))
    collect(ClaudeCvParser(client=client, model="m").parse_pdf(b"%PDF-1.4"))
    block = messages.kwargs["messages"][0]["content"][0]
    assert block["type"] == "document" and block["source"]["media_type"] == "application/pdf"


@pytest.mark.parametrize("stop_reason,parsed", [("refusal", None), ("max_tokens", None), ("end_turn", None)])
def test_claude_parser_unusable_responses_fail_each_part(stop_reason, parsed):
    client, _ = fake_client(SimpleNamespace(stop_reason=stop_reason, parsed_output=parsed, _request_id="r"))
    resume, errors = collect(ClaudeCvParser(client=client, model="m").parse_text("x"))
    assert resume == {} and set(errors) == set(PARTS)


def test_claude_parser_cut_off_json_fails_only_that_part():
    output = ParsedResume.model_validate({"basics": {"name": "Jane Doe"}})
    client, messages = fake_client(SimpleNamespace(stop_reason="end_turn", parsed_output=output, _request_id="r"))
    real_parse = messages.parse

    def parse(**kwargs):
        if kwargs["output_format"] is ExtraSections:  # the SDK raises when the JSON is cut off
            ExtraSections.model_validate_json('{"skills": [{"name": "SQ')
        return real_parse(**kwargs)

    messages.parse = parse
    resume, errors = collect(ClaudeCvParser(client=client, model="m").parse_text("x"))
    assert set(errors) == {"extras"} and resume["basics"]["name"] == "Jane Doe"


def test_cv_row_defaults(client, auth):
    client.post("/api/cv/upload", headers=auth[0], files={"file": ("cv.pdf", make_pdf(), PDF_MIME)})
    with client.session_factory() as db:
        cv = db.query(CV).one()
        assert cv.json_resume == {} and cv.parsing_error is None


def test_current_role_is_explicit_in_schema_and_prompt():
    from app.services.cv_parser import EXPERIENCE_TASK

    work = ParsedResume.model_validate({"basics": {}, "work": [{"name": "Acme"}]}).work[0]
    assert work.current is False  # a missing end date alone does not mean "current"
    assert "current to true" in EXPERIENCE_TASK


def test_parser_without_api_key_fails_each_part_with_a_clear_log(monkeypatch, caplog):
    import anthropic

    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_AUTH_TOKEN", raising=False)
    resume, errors = collect(ClaudeCvParser(client=anthropic.Anthropic(api_key=None), model="m").parse_text("x"))
    assert resume == {} and set(errors) == set(PARTS)
    assert "not available right now" in str(next(iter(errors.values())))
    assert "No Claude API key configured" in caplog.text


def test_parser_client_uses_cvsmith_key_and_pinned_base_url(monkeypatch):
    from app.config import get_settings

    monkeypatch.setenv("CVSMITH_CLAUDE_KEY", "sk-test-cvsmith")
    monkeypatch.setenv("ANTHROPIC_BASE_URL", "http://host-internal.invalid")
    get_settings.cache_clear()
    try:
        client = ClaudeCvParser().client
        assert client.api_key == "sk-test-cvsmith"
        assert str(client.base_url).rstrip("/") == "https://api.anthropic.com"
    finally:
        get_settings.cache_clear()


def test_schema_has_no_optional_unions():
    # The API compiles the output schema into a grammar; ~30 Optional fields (each an anyOf)
    # made it "too large" and every real parse failed with a 400.
    from anthropic.lib._parse._transform import transform_schema

    for schema in (BasicsPart, ExperiencePart, ExtraSections):
        assert "anyOf" not in json.dumps(transform_schema(schema))


def test_each_request_schema_stays_well_under_the_grammar_limit():
    # The full schema with every section was rejected as "too large"; each half must stay
    # small. ~5.5k chars was accepted, so keep each request's schema below that.
    from anthropic.lib._parse._transform import transform_schema

    for schema in (BasicsPart, ExperiencePart, ExtraSections):
        assert len(json.dumps(transform_schema(schema))) < 5000, schema.__name__


def test_parser_combines_both_requests_including_fresher_sections():
    output = ParsedResume.model_validate({
        "basics": {"name": "Jane"},
        "work": [{"name": "Acme"}],
        "languages": [{"language": "Tamil", "fluency": "Native"}],
        "awards": [{"title": "Employee of the month", "date": "2018-08"}],
        "volunteer": [{"organization": "College symposium", "position": "Sponsorship head"}],
        "projects": [{"name": "Portfolio site"}],
        "interests": [{"name": "Chess"}],
    })
    client, _ = fake_client(SimpleNamespace(stop_reason="end_turn", parsed_output=output, _request_id="r"))
    result, _ = collect(ClaudeCvParser(client=client, model="m").parse_text("x"))
    assert result["basics"]["name"] == "Jane" and result["work"][0]["name"] == "Acme"
    assert result["languages"] == [{"language": "Tamil", "fluency": "Native"}]
    assert result["awards"] == [{"title": "Employee of the month", "date": "2018-08"}]
    assert result["volunteer"][0]["position"] == "Sponsorship head"
    assert result["projects"] == [{"name": "Portfolio site", "highlights": []}]
    assert result["interests"] == [{"name": "Chess", "keywords": []}]


def test_one_failed_request_does_not_lose_the_other_parts():
    class OneFails(FakeMessages):
        def parse(self, **kwargs):
            if kwargs["output_format"] is ExtraSections:
                return SimpleNamespace(stop_reason="refusal", parsed_output=None, _request_id="r")
            return super().parse(**kwargs)

    output = ParsedResume.model_validate({"basics": {"name": "Jane"}, "skills": [{"name": "SQL"}]})
    messages = OneFails(SimpleNamespace(stop_reason="end_turn", parsed_output=output, _request_id="r"))
    client = SimpleNamespace(beta=SimpleNamespace(messages=messages))
    resume, errors = collect(ClaudeCvParser(client=client, model="m").parse_text("x"))
    assert set(errors) == {"extras"} and resume["basics"]["name"] == "Jane" and "skills" not in resume


def test_summary_from_the_experience_part_joins_the_contact_details():
    output = ParsedResume.model_validate({"basics": {"name": "Jane", "summary": "Analyst."}})
    client, _ = fake_client(SimpleNamespace(stop_reason="end_turn", parsed_output=output, _request_id="r"))
    result, _ = collect(ClaudeCvParser(client=client, model="m").parse_text("x"))
    assert result["basics"] == {"name": "Jane", "summary": "Analyst.", "profiles": []}


def test_parser_drops_empty_values_but_keeps_lists():
    output = ParsedResume.model_validate({"basics": {"name": "Jane"}, "work": [{"name": "Acme", "current": True}]})
    client, _ = fake_client(SimpleNamespace(stop_reason="end_turn", parsed_output=output, _request_id="r"))
    result, _ = collect(ClaudeCvParser(client=client, model="m").parse_text("x"))
    assert result["basics"] == {"name": "Jane", "profiles": []}
    assert result["work"] == [{"name": "Acme", "current": True, "highlights": []}]


def test_extras_part_does_not_repeat_job_achievements():
    # A real scan put "Employee of the month" both under the job and in awards.
    from app.services.cv_parser import EXTRA_TASK

    assert "never repeat it here" in EXTRA_TASK and "outside any job" in EXTRA_TASK


def test_contact_patterns_take_only_unambiguous_details():
    from app.services.contact_patterns import find_contact

    text = "Jane Doe\nJane.Doe@Example.com | +44 20 7946 0958 | linkedin.com/in/jane-doe/\nAcme 2014 - 2018"
    assert find_contact(text) == {"basics": {
        "email": "Jane.Doe@Example.com", "phone": "+44 20 7946 0958",
        "profiles": [{"network": "LinkedIn", "url": "linkedin.com/in/jane-doe"}],
    }}
    # year ranges, dates and long reference numbers are not phone numbers
    assert find_contact("Worked 2014 - 2018, 2012-2014; 06/2014 - 05/2018; ID 1234567890123456789") == {}
    assert find_contact("") == {}


def test_contact_step_never_decides_success():
    # the instant contact step always completes; whether the CV was read depends on the AI parts
    from app.models.cv import AI_PARTS

    assert "contact" not in AI_PARTS


def test_uploaded_file_is_deleted_once_parsing_has_finished(client, stored_cv):
    from pathlib import Path

    def files():
        return [p for p in Path(client.upload_dir).rglob("*") if p.is_file()]

    body, _ = stored_cv(make_text_pdf(CV_TEXT))
    assert body["parsing_status"] == "complete" and files() == []
    # also when nothing could be read
    body, _ = stored_cv(make_pdf())
    assert body["parsing_status"] == "failed" and files() == []


def make_docx_with_header(body, header):
    document = docx.Document()
    for p in body:
        document.add_paragraph(p)
    document.sections[0].header.paragraphs[0].text = header
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_contact_details_in_a_word_header_are_found(stored_cv):
    data = make_docx_with_header(CV_TEXT.split("\n")[1:], "Jane Doe | header@example.com | +44 7700 900123")
    body, _ = stored_cv(data, "cv.docx", DOCX_MIME)
    basics = body["json_resume"]["basics"]
    assert basics["email"] == "header@example.com" and basics["phone"] == "+44 7700 900123"


def test_header_text_is_not_sent_to_the_ai(stored_cv):
    data = make_docx_with_header(CV_TEXT.split("\n"), "Secret header line")
    _, parser = stored_cv(data, "cv.docx", DOCX_MIME)
    assert "Secret header line" not in parser.calls[0][1]
